"""時光機＋算命仙的多語言字串表：抽字、匯出 Excel、匯回、產生 Mod 用的 strings.tsv。

key 是繁中原文；「繁中（修改）」有填就用修改後的句子，日文／韓文有填就用翻譯，沒填的顯示原文。
主檔是 i18n/strings.json（進 git）；Excel 只是編輯用（不進 git，隨時可重新匯出）。

用法（在 repo 根目錄）：
  python -I tools/i18n.py extract           從網頁範本、guides.json、data/graph.json（其他條件）、時光機 C# 抽字，合併進 strings.json
  python -I tools/i18n.py export [--force]  strings.json → i18n/strings.xlsx（Excel 比 strings.json 新時要 --force 才覆蓋）
  python -I tools/i18n.py import            i18n/strings.xlsx → strings.json（檢查 {0} 參數有沒有漏）
  python -I tools/i18n.py tsv <輸出路徑>     產生時光機 Mod 讀的 strings.tsv
"""
import glob
import html.parser
import json
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MASTER = os.path.join(ROOT, "i18n", "strings.json")
XLSX = os.path.join(ROOT, "i18n", "strings.xlsx")
TEMPLATE = os.path.join(ROOT, "tools", "flowchart_template.html")
GUIDES = os.path.join(ROOT, "tools", "guides.json")
GRAPH = os.path.join(ROOT, "data", "graph.json")
LANGS = ["zh-TW", "ja", "ko"]
CJK = re.compile(r"[\u4e00-\u9fff]")
# 分頁：key 是 strings.json 的 group
SHEETS = [("web", "算命仙網頁"), ("guide", "結緣與其他條件"), ("mod", "時光機視窗")]


# ---------- 抽字 ----------
def scan_args(src, i):
    """src[i] 是 T( 的第一個參數之後：讀到對應的右括號，回傳各參數的原始碼（給「變數」欄看 {0}{1} 是什麼）"""
    args, depth, start, j = [], 0, i, i
    while j < len(src):
        c = src[j]
        if c in "\"'`":
            j = skip_string(src, j)
            continue
        if c in "([{":
            depth += 1
        elif c in ")]}":
            if depth == 0:
                break
            depth -= 1
        elif c == "," and depth == 0:
            args.append(src[start:j])
            start = j + 1
        j += 1
    args.append(src[start:j])
    return [re.sub(r"\s+", " ", a).strip() for a in args[1:]]  # 第一段是字串後面到第一個逗號之間（空的）


def skip_string(src, j):
    q = src[j]
    j += 1
    while src[j] != q:
        if src[j] == "\\":
            j += 1
        elif q == "`" and src.startswith("${", j):
            depth, j = 1, j + 2
            while depth:
                if src[j] in "\"'`":
                    j = skip_string(src, j)
                    continue
                depth += {"{": 1, "}": -1}.get(src[j], 0)
                j += 1
            continue
        j += 1
    return j + 1


def line_of(src, pos):
    return src.count("\n", 0, pos) + 1


def extract_web(found):
    page = open(TEMPLATE, encoding="utf-8").read()
    start, end = page.rindex("<script>") + 8, page.rindex("</script>")
    script = page[start:end]
    base = line_of(page, start) - 1
    for m in re.finditer(r'\bT\(("(?:[^"\\\n]|\\.)*")', script):
        key = json.loads(m.group(1))
        args = scan_args(script, m.end())
        add(found, key, "web", f"網頁 {base + line_of(script, m.start())} 行", args)

    # 固定的 HTML（分頁名稱、勾選項、提示）：網頁載入時 i18nStatic 會換掉
    class P(html.parser.HTMLParser):
        skip = 1  # <body> 之前（網頁只換 body 裡的字）

        def handle_starttag(self, tag, attrs):
            if tag == "body":
                self.skip -= 1
            if tag in ("script", "style"):
                self.skip += 1
            for k, v in attrs:
                if k in ("title", "placeholder") and v and CJK.search(v) and not self.skip:
                    add(found, v, "web", f"網頁 {self.getpos()[0]} 行（{k}）", [])

        def handle_endtag(self, tag):
            if tag in ("script", "style"):
                self.skip -= 1

        def handle_data(self, data):
            t = data.strip()
            if not self.skip and t and CJK.search(t):
                add(found, t, "web", f"網頁 {self.getpos()[0]} 行", [])

    P(convert_charrefs=True).feed(page[:start])


def extract_guides(found):
    g = json.load(open(GUIDES, encoding="utf-8"))
    for kind, label in (("romance", "結緣條件"), ("targets", "其他條件")):
        for guide in g.get(kind, []):
            where = f"{label}：{guide['name']}"
            if guide.get("intro"):
                add(found, guide["intro"], "guide", where, [])

            def items(lst):
                for it in lst:
                    add(found, it["text"], "guide", where, [])
                    items(it.get("sub", []))
            items(guide["items"])
    # 稱號列表的取得方式
    for name, text in g.get("titles", {}).items():
        add(found, text, "guide", f"稱號列表：{name}", [])
    # 其他條件的精簡版（build_graph.py 的 TARGETS）：從產生好的 graph.json 抽
    if not os.path.exists(GRAPH):
        print("（沒有 data/graph.json，略過其他條件的 TARGETS 文字；先跑 rebuild_guide.sh）")
        return
    for t in json.load(open(GRAPH, encoding="utf-8"))["targets"]:
        where = f"其他條件：{t['name']}"
        add(found, t["name"], "guide", where, [])
        if t.get("note"):
            add(found, t["note"], "guide", where, [])
        for c in (t.get("conds") or []) + (t.get("steps") or []):
            if isinstance(c, str):
                add(found, c, "guide", where, [])
            else:
                add(found, c["t"], "guide", where, [])
                if c.get("pop"):
                    add(found, c["pop"]["text"], "guide", where + "（可點的那段字，要出現在上一句裡）", [])


def cs_unescape(s):
    return re.sub(r'\\(.)', lambda m: {"n": "\n", "t": "\t"}.get(m.group(1), m.group(1)), s)


def extract_cs(found):
    for p in sorted(glob.glob(os.path.join(ROOT, "TimeMachine", "*.cs"))):
        src = open(p, encoding="utf-8-sig").read()
        for m in re.finditer(r'\bL\.T\("((?:[^"\\\n]|\\.)*)"', src):
            add(found, cs_unescape(m.group(1)), "mod", f"{os.path.basename(p)} {line_of(src, m.start())} 行", scan_args(src, m.end()))


def add(found, key, group, where, args):
    e = found.setdefault(key, {"groups": [], "where": [], "vars": []})
    if group not in e["groups"]:
        e["groups"].append(group)
    if where not in e["where"]:
        e["where"].append(where)
    if args and not e["vars"]:
        e["vars"] = args


def load_master():
    if not os.path.exists(MASTER):
        return {"_說明": __doc__.strip().split("\n")[0], "strings": {}}
    return json.load(open(MASTER, encoding="utf-8"))


def save_master(m):
    os.makedirs(os.path.dirname(MASTER), exist_ok=True)
    with open(MASTER, "w", encoding="utf-8", newline="\n") as f:
        json.dump(m, f, ensure_ascii=False, indent=1)
        f.write("\n")


def cmd_extract():
    found = {}
    extract_web(found)
    extract_guides(found)
    extract_cs(found)
    m = load_master()
    old = m["strings"]
    new, dropped = {}, []
    for key, e in found.items():
        prev = old.get(key, {})
        new[key] = {**{lang: prev.get(lang, "") for lang in LANGS}, "note": prev.get("note", ""),
                    "group": e["groups"][0], "also": e["groups"][1:], "where": e["where"], "vars": e["vars"]}
    # 原文已經不在程式裡的：有改寫或翻譯的留著（標「不再使用」），空的直接拿掉
    for key, prev in old.items():
        if key in new:
            continue
        if any(prev.get(lang) for lang in LANGS) or prev.get("note"):
            new[key] = {**prev, "unused": True}
        else:
            dropped.append(key)
    m["strings"] = new
    save_master(m)
    counts = {g: sum(1 for e in new.values() if e["group"] == g and not e.get("unused")) for g, _ in SHEETS}
    print("strings.json：", "、".join(f"{label} {counts[g]}" for g, label in SHEETS),
          f"；不再使用但有內容 {sum(1 for e in new.values() if e.get('unused'))}；移除 {len(dropped)}")


# ---------- Excel ----------
# Excel 存不了控制字元（網頁有用 \u0000 當記號）：寫成 \u0000 這樣的字，匯回時還原
vis = lambda t: re.sub(r"[\x00-\x08\x0b-\x1f]", lambda m: "\\u%04x" % ord(m.group()), t)
unvis = lambda t: re.sub(r"\\u(000[0-8]|000[bcdef]|001[0-9a-f])", lambda m: chr(int(m.group(1), 16)), t)
HEAD = ["原文（不要改）", "繁中（修改）", "日文", "韓文", "備註", "變數", "位置"]


def cmd_export(force):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    if os.path.exists(XLSX) and os.path.getmtime(XLSX) > os.path.getmtime(MASTER) and not force:
        sys.exit("i18n/strings.xlsx 比 strings.json 新，可能有還沒匯回的修改。先跑 import，或加 --force 覆蓋。")
    m = load_master()
    wb = Workbook()
    wb.remove(wb.active)
    guide = wb.create_sheet("說明")
    for row in [
        ["這份表是時光機（F6）與算命仙網頁（F7）的介面文字。"],
        ["「原文」是程式裡的 key，不要改；要改說法填「繁中（修改）」，翻譯填「日文」「韓文」。沒填的格子顯示原文。"],
        ["{0}、{1}… 是程式填進去的值（看「變數」欄），改寫或翻譯時要保留，順序可以換。"],
        ["<b>、<span …> 這類 HTML 標籤要保留成對；\\n 是換行。"],
        ["「結緣與其他條件」分頁的文字也可以直接改 tools/guides.json、build_graph.py 的 TARGETS；改原文後要重新 extract，舊的修改會變成「不再使用」。"],
        ["改完存檔後跑：python -I tools/i18n.py import，再跑 tools/rebuild_guide.sh。"],
    ]:
        guide.append(row)
    guide.column_dimensions["A"].width = 120
    fill = PatternFill("solid", fgColor="EEEEEE")
    for group, label in SHEETS + [("unused", "不再使用")]:
        rows = [(k, e) for k, e in m["strings"].items() if (e.get("unused") if group == "unused" else e["group"] == group and not e.get("unused"))]
        if not rows:
            continue
        ws = wb.create_sheet(label)
        ws.append(HEAD)
        for k, e in rows:
            where = "、".join(e.get("where", []))
            if e.get("also"):
                where += "（也用在：" + "、".join(dict(SHEETS)[g] for g in e["also"]) + "）"
            ws.append([vis(k)] + [vis(e.get(x, "")) for x in LANGS + ["note"]] + [
                       "\n".join(f"{{{i}}} = {v}" for i, v in enumerate(e.get("vars", []))), where])
        for col, w in zip("ABCDEFG", [60, 60, 40, 40, 24, 36, 30]):
            ws.column_dimensions[col].width = w
        for r in ws.iter_rows():
            for c in r:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in ws.iter_rows(min_row=2, max_col=1):
            r[0].fill = fill
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(XLSX)
    print("已匯出", XLSX)


def cmd_import():
    from openpyxl import load_workbook
    m = load_master()
    s = m["strings"]
    changed, problems = 0, []
    for ws in load_workbook(XLSX).worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows or list(rows[0][:len(HEAD)]) != HEAD:
            continue
        for r in rows[1:]:
            key = unvis(r[0] or "")
            if key not in s:
                if key:
                    problems.append(f"找不到原文（是不是改到原文欄了？）：{key[:40]}")
                continue
            vals = dict(zip(LANGS + ["note"], [unvis(v or "") if isinstance(v, str) or v is None else str(v) for v in r[1:5]]))
            if vals["zh-TW"] == key:
                vals["zh-TW"] = ""
            for lang in LANGS:
                if vals[lang]:
                    need = set(re.findall(r"\{\d+[^}]*\}", key))
                    have = set(re.findall(r"\{\d+[^}]*\}", vals[lang]))
                    if need != have:
                        problems.append(f"[{ws.title}] {lang} 的參數和原文不同（原文 {sorted(need)}，填的 {sorted(have)}）：{key[:40]}")
                    if len(re.findall(r"<[a-z]", key)) != len(re.findall(r"<[a-z]", vals[lang])):
                        problems.append(f"[{ws.title}] {lang} 的 HTML 標籤數和原文不同：{key[:40]}")
            for k, v in vals.items():
                if s[key].get(k, "") != v:
                    s[key][k] = v
                    changed += 1
    save_master(m)
    print(f"匯回 {changed} 格")
    for p in problems:
        print("⚠", p)


def cmd_tsv(out):
    esc = lambda t: t.replace("\\", "\\\\").replace("\t", "\\t").replace("\n", "\\n")
    m = load_master()
    lines = ["key\t" + "\t".join(LANGS)]
    for k, e in m["strings"].items():
        if (e["group"] == "mod" or "mod" in e.get("also", [])) and any(e.get(lang) for lang in LANGS):
            lines.append("\t".join([esc(k)] + [esc(e.get(lang, "")) for lang in LANGS]))
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    with open(out, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines) + "\n")
    print("已寫出", out, len(lines) - 1, "筆")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "extract":
        cmd_extract()
    elif cmd == "export":
        cmd_export("--force" in sys.argv)
    elif cmd == "import":
        cmd_import()
    elif cmd == "tsv" and len(sys.argv) > 2:
        cmd_tsv(sys.argv[2])
    else:
        print(__doc__)
